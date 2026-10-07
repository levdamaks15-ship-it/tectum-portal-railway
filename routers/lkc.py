import io
import os
from datetime import datetime, date, timedelta
from typing import Optional, List, Dict, Any
from fastapi import APIRouter, Depends, HTTPException, Request, Response, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, selectinload
from sqlalchemy import desc, asc, and_
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import models
import schemas
from routers.common import get_db

router = APIRouter(prefix="/api/lkc", tags=["lkc"])

# --- СПРАВОЧНИК НОМЕНКЛАТУРЫ И НОРМ РАСХОДА ---
LKC_CATALOG = [
    {
        "id": "strip_1800_300_6",
        "name": "Полоса плоская 6мм 1800х300",
        "product_type": "Плоская полоса",
        "dimension_size": "пл. лист 6мм 1800х300",
        "default_unit": "шт",
        "norm_t": 0,
        "norm_l": 0,
        "norm_screws": 0,
        "plan_default": 750,
        "description": "Нарезка полос из листа 6мм 1800х1200 (4 полосы с 1 листа)"
    },
    {
        "id": "strip_1800_300_8",
        "name": "Полоса плоская 8мм 1800х300",
        "product_type": "Плоская полоса",
        "dimension_size": "пл. лист 8мм 1800х300",
        "default_unit": "шт",
        "norm_t": 0,
        "norm_l": 0,
        "norm_screws": 0,
        "plan_default": 750,
        "description": "Нарезка полос из листа 8мм 1800х1200 (4 полосы с 1 листа)"
    },
    {
        "id": "sheet_1800_1200_6",
        "name": "Плоский лист 6мм 1800х1200",
        "product_type": "Плоский лист",
        "dimension_size": "пл. лист 6мм 1800х1200",
        "default_unit": "шт",
        "norm_t": 0,
        "norm_l": 0,
        "norm_screws": 0,
        "plan_default": 750,
        "description": "Плоский хризотилцементный лист 6 мм. Сменный план: 750 шт"
    },
    {
        "id": "sheet_1800_1200_8",
        "name": "Плоский лист 8мм 1800х1200",
        "product_type": "Плоский лист",
        "dimension_size": "пл. лист 8мм 1800х1200",
        "default_unit": "шт",
        "norm_t": 0,
        "norm_l": 0,
        "norm_screws": 0,
        "plan_default": 750,
        "description": "Базовый плоский хризотилцементный лист 8 мм. Сменный план: 750 шт"
    },
    {
        "id": "bed_3600",
        "name": "Грядка 3.6 м (3600х900)",
        "product_type": "Комплект грядок",
        "dimension_size": "грядка 3600х900",
        "default_unit": "компл",
        "norm_t": 2,
        "norm_l": 4,
        "norm_screws": 24,
        "plan_default": 750,
        "description": "2 продольных борта (1800+1800) + 2 торца 900х245. Расход: Т-профиль 2 шт, L-уголок 4 шт, Шурупы 24 шт"
    },
    {
        "id": "bed_2700",
        "name": "Грядка 2.7 м (2700х900)",
        "product_type": "Комплект грядок",
        "dimension_size": "грядка 2700х900",
        "default_unit": "компл",
        "norm_t": 2,
        "norm_l": 4,
        "norm_screws": 24,
        "plan_default": 750,
        "description": "2 борта (1800+900) + 2 торца 900х245. Расход: Т-профиль 2 шт, L-уголок 4 шт, Шурупы 24 шт"
    },
    {
        "id": "bed_1800",
        "name": "Грядка 1.8 м (1800х900)",
        "product_type": "Комплект грядок",
        "dimension_size": "грядка 1800х900",
        "default_unit": "компл",
        "norm_t": 0,
        "norm_l": 4,
        "norm_screws": 16,
        "plan_default": 750,
        "description": "2 борта 1800х245 + 2 торца 900х245. Расход: Т-профиль 0 шт, L-уголок 4 шт, Шурупы 16 шт"
    }
]

@router.get("/catalog")
def get_catalog():
    """Возвращает справочник типоразмеров и норм расхода комплектующих."""
    return {"catalog": LKC_CATALOG}


@router.get("/reports", response_model=List[schemas.LKCShiftReportOut])
def get_reports(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    limit: int = 50,
    db: Session = Depends(get_db)
):
    """Возвращает список рапортов ЛКЦ с фильтрацией по датам."""
    query = db.query(models.LKCShiftReport).options(
        selectinload(models.LKCShiftReport.items),
        selectinload(models.LKCShiftReport.raw_materials),
        selectinload(models.LKCShiftReport.downtimes)
    ).filter(models.LKCShiftReport.is_active == True)

    if start_date:
        query = query.filter(models.LKCShiftReport.report_date >= start_date)
    if end_date:
        query = query.filter(models.LKCShiftReport.report_date <= end_date)

    reports = query.order_by(models.LKCShiftReport.report_date.desc(), models.LKCShiftReport.id.desc()).limit(limit).all()
    return reports


@router.get("/reports/by_date", response_model=Optional[schemas.LKCShiftReportOut])
def get_report_by_date(
    report_date: date,
    shift_name: str = "Дневная смена (08:00 - 17:00)",
    db: Session = Depends(get_db)
):
    """Получает рапорт ЛКЦ на конкретную дату и смену."""
    report = db.query(models.LKCShiftReport).options(
        selectinload(models.LKCShiftReport.items),
        selectinload(models.LKCShiftReport.raw_materials),
        selectinload(models.LKCShiftReport.downtimes)
    ).filter(
        models.LKCShiftReport.report_date == report_date,
        models.LKCShiftReport.shift_name == shift_name,
        models.LKCShiftReport.is_active == True
    ).first()
    return report


@router.get("/reports/{report_id}", response_model=schemas.LKCShiftReportOut)
def get_report_by_id(report_id: int, db: Session = Depends(get_db)):
    """Получает подробный рапорт ЛКЦ по ID."""
    report = db.query(models.LKCShiftReport).options(
        selectinload(models.LKCShiftReport.items),
        selectinload(models.LKCShiftReport.raw_materials),
        selectinload(models.LKCShiftReport.downtimes)
    ).filter(
        models.LKCShiftReport.id == report_id,
        models.LKCShiftReport.is_active == True
    ).first()
    if not report:
        raise HTTPException(status_code=404, detail="Рапорт ЛКЦ не найден")
    return report


@router.post("/reports", response_model=schemas.LKCShiftReportOut)
def save_report(
    data: schemas.LKCShiftReportCreate,
    request: Request,
    db: Session = Depends(get_db)
):
    """
    Создает или обновляет рапорт смены ЛКЦ в единой атомарной транзакции (Strict Backend-First).
    """
    try:
        # Ищем существующий рапорт на эту дату и смену
        report = db.query(models.LKCShiftReport).filter(
            models.LKCShiftReport.report_date == data.report_date,
            models.LKCShiftReport.shift_name == data.shift_name,
            models.LKCShiftReport.is_active == True
        ).first()

        if not report:
            report = models.LKCShiftReport(
                report_date=data.report_date,
                shift_name=data.shift_name,
                master_id=data.master_id,
                master_name=data.master_name or "Миркасимов И.М.",
                status=data.status or "completed",
                raw_sheets_thickness=data.raw_sheets_thickness or "8 мм",
                raw_sheets_spent=data.raw_sheets_spent or 0.0,
                warehouse_submitted_qty=data.warehouse_submitted_qty or 0.0,
                scrap_defect_qty=data.scrap_defect_qty or 0.0,
                notes=data.notes or "",
                is_active=True,
                created_at=datetime.utcnow(),
                updated_at=datetime.utcnow()
            )
            db.add(report)
            db.flush()
        else:
            # Обновляем поля шапки
            report.master_id = data.master_id
            report.master_name = data.master_name or report.master_name
            report.status = data.status or report.status
            report.raw_sheets_thickness = data.raw_sheets_thickness or report.raw_sheets_thickness
            report.raw_sheets_spent = data.raw_sheets_spent or 0.0
            report.warehouse_submitted_qty = data.warehouse_submitted_qty or 0.0
            report.scrap_defect_qty = data.scrap_defect_qty or 0.0
            report.notes = data.notes or ""
            report.updated_at = datetime.utcnow()
            
            # Удаляем старые строки позиций для чистой перезаписи
            db.query(models.LKCReportItem).filter(models.LKCReportItem.report_id == report.id).delete()
            db.query(models.LKCRawMaterialUsage).filter(models.LKCRawMaterialUsage.report_id == report.id).delete()

        # Добавляем строки позиций выпуска
        for idx, item_data in enumerate(data.items or []):
            item = models.LKCReportItem(
                report_id=report.id,
                item_order=idx + 1,
                product_type=item_data.product_type or "Комплект грядок",
                dimension_size=item_data.dimension_size or "",
                plan_qty=item_data.plan_qty or 0.0,
                fact_qty=item_data.fact_qty or 0.0,
                fact_unit=item_data.fact_unit or "шт",
                order_batch=item_data.order_batch or "",
                sets_count=item_data.sets_count or 0,
                comp_t_count=item_data.comp_t_count or 0.0,
                comp_l_count=item_data.comp_l_count or 0.0,
                comp_screws_count=item_data.comp_screws_count or 0,
                raw_sheets_thickness=item_data.raw_sheets_thickness or report.raw_sheets_thickness,
                raw_sheets_spent=item_data.raw_sheets_spent or 0.0,
                warehouse_submitted_qty=item_data.warehouse_submitted_qty or 0.0,
                notes=item_data.notes or "",
                created_at=datetime.utcnow()
            )
            db.add(item)

        # Добавляем строки списания сырья
        total_raw_spent = 0.0
        total_wh_sub = 0.0
        total_scrap = 0.0
        for idx, raw_data in enumerate(data.raw_materials or []):
            spent = raw_data.spent_qty or 0.0
            wh_qty = raw_data.warehouse_submitted_qty or 0.0
            scrap = raw_data.scrap_defect_qty or 0.0
            total_raw_spent += spent
            total_wh_sub += wh_qty
            total_scrap += scrap

            raw_obj = models.LKCRawMaterialUsage(
                report_id=report.id,
                item_order=idx + 1,
                material_name=raw_data.material_name or "Плоский лист 8мм 1800х1200",
                thickness=raw_data.thickness or "8 мм",
                spent_qty=spent,
                unit=raw_data.unit or "шт",
                warehouse_submitted_qty=wh_qty,
                scrap_defect_qty=scrap,
                notes=raw_data.notes or "",
                created_at=datetime.utcnow()
            )
            db.add(raw_obj)

        if data.raw_materials:
            report.raw_sheets_spent = total_raw_spent
            report.warehouse_submitted_qty = total_wh_sub
            report.scrap_defect_qty = total_scrap

        # Обработка простоев (если переданы)
        if data.downtimes:
            # Удаляем старые простои, привязанные к отчету
            db.query(models.LKCDowntime).filter(models.LKCDowntime.report_id == report.id).delete()
            for dt in data.downtimes:
                downtime_obj = models.LKCDowntime(
                    report_id=report.id,
                    downtime_date=data.report_date,
                    shift_name=data.shift_name,
                    master_name=report.master_name,
                    start_time=dt.start_time,
                    end_time=dt.end_time or "",
                    duration_minutes=dt.duration_minutes or 0,
                    equipment_node=dt.equipment_node or "Дисковая пила",
                    category=dt.category or "Механическая",
                    reason=dt.reason or "",
                    comment=dt.comment or "",
                    is_active=True,
                    created_at=datetime.utcnow()
                )
                db.add(downtime_obj)

        db.commit()

        # Перезагружаем со всеми связями
        db.refresh(report)
        return db.query(models.LKCShiftReport).options(
            selectinload(models.LKCShiftReport.items),
            selectinload(models.LKCShiftReport.raw_materials),
            selectinload(models.LKCShiftReport.downtimes)
        ).filter(models.LKCShiftReport.id == report.id).first()

    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Ошибка сохранения рапорта ЛКЦ: {str(e)}")


@router.delete("/reports/{report_id}")
def delete_report(report_id: int, db: Session = Depends(get_db)):
    """Мягкое удаление рапорта ЛКЦ."""
    report = db.query(models.LKCShiftReport).filter(models.LKCShiftReport.id == report_id).first()
    if not report:
        raise HTTPException(status_code=404, detail="Рапорт не найден")
    report.is_active = False
    report.updated_at = datetime.utcnow()
    db.commit()
    return {"status": "ok", "message": f"Рапорт №{report_id} успешно удален"}


# --- ПРОСТОИ ЛКЦ ---
@router.get("/downtimes", response_model=List[schemas.LKCDowntimeOut])
def get_downtimes(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db)
):
    """Возвращает список простоев участка ЛКЦ."""
    query = db.query(models.LKCDowntime).filter(models.LKCDowntime.is_active == True)
    if start_date:
        query = query.filter(models.LKCDowntime.downtime_date >= start_date)
    if end_date:
        query = query.filter(models.LKCDowntime.downtime_date <= end_date)
    return query.order_by(models.LKCDowntime.downtime_date.desc(), models.LKCDowntime.start_time.desc()).all()


@router.post("/downtimes", response_model=schemas.LKCDowntimeOut)
def create_downtime(data: schemas.LKCDowntimeCreate, db: Session = Depends(get_db)):
    """Создает запись простоя на участке ЛКЦ."""
    try:
        dt = models.LKCDowntime(
            report_id=data.report_id,
            downtime_date=data.downtime_date or date.today(),
            shift_name=data.shift_name or "Дневная смена",
            master_name=data.master_name or "Миркасимов И.М.",
            start_time=data.start_time,
            end_time=data.end_time or "",
            duration_minutes=data.duration_minutes or 0,
            equipment_node=data.equipment_node or "Дисковая пила",
            category=data.category or "Механическая",
            reason=data.reason or "",
            comment=data.comment or "",
            is_active=True,
            created_at=datetime.utcnow()
        )
        db.add(dt)
        db.commit()
        db.refresh(dt)
        return dt
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Ошибка создания записи простоя: {str(e)}")


@router.delete("/downtimes/{downtime_id}")
def delete_downtime(downtime_id: int, db: Session = Depends(get_db)):
    """Мягкое удаление записи простоя."""
    dt = db.query(models.LKCDowntime).filter(models.LKCDowntime.id == downtime_id).first()
    if not dt:
        raise HTTPException(status_code=404, detail="Запись простоя не найдена")
    dt.is_active = False
    db.commit()
    return {"status": "ok", "message": f"Простой №{downtime_id} удален"}


# --- ЭКСПОРТ В EXCEL ---
@router.get("/export_excel")
def export_excel(
    start_date: Optional[date] = None,
    end_date: Optional[date] = None,
    db: Session = Depends(get_db)
):
    """
    Генерирует недельный рапорт ЛКЦ в формате Excel (.xlsx),
    в точности повторяющий официальный шаблон предприятия.
    """
    if not start_date:
        today = date.today()
        start_date = today - timedelta(days=today.weekday())
    if not end_date:
        end_date = start_date + timedelta(days=6)

    # Извлекаем все рапорты за период
    reports = db.query(models.LKCShiftReport).options(
        selectinload(models.LKCShiftReport.items),
        selectinload(models.LKCShiftReport.downtimes)
    ).filter(
        models.LKCShiftReport.report_date >= start_date,
        models.LKCShiftReport.report_date <= end_date,
        models.LKCShiftReport.is_active == True
    ).order_by(models.LKCShiftReport.report_date.asc(), models.LKCShiftReport.id.asc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Недельный рапорт ЛКЦ"

    # Стили
    title_font = Font(name="Calibri", size=13, bold=True, color="1E3A8A")
    sub_font = Font(name="Calibri", size=9, bold=False, color="475569")
    header_font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=9, bold=False, color="0F172A")
    bold_font = Font(name="Calibri", size=9, bold=True, color="0F172A")
    total_font = Font(name="Calibri", size=10, bold=True, color="1E3A8A")

    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    sub_header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    total_fill = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )
    thick_bottom = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="medium", color="1E3A8A")
    )

    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center")

    # Шапка отчета
    ws.merge_cells("A1:M1")
    ws["A1"] = "ТЕКТУМ ИНЖИНИРИНГ (TECTUM ENGINEERING) — НЕДЕЛЬНЫЙ РАПОРТ ЛКЦ"
    ws["A1"].font = title_font
    ws["A1"].alignment = align_center

    period_str = f"Период: {start_date.strftime('%d.%m.%Y')} – {end_date.strftime('%d.%m.%Y')} | Участок: Линия кровельных сэндвич-панелей и комплектующих (ЛКЦ) | Режим: Дневная смена"
    ws.merge_cells("A2:M2")
    ws["A2"] = period_str
    ws["A2"].font = sub_font
    ws["A2"].alignment = align_center

    # Заголовки таблицы (Строки 4 и 5)
    headers = [
        ("A4:A5", "№\nп/п"),
        ("B4:B5", "Дата"),
        ("C4:C5", "Мастер\n(Ф.И.О.)"),
        ("D4:D5", "Размер /\nТипоразмер"),
        ("E4:E5", "План\n(шт.)"),
        ("F4:F5", "Факт\n(шт.)"),
        ("G4:G5", "Комплект\n(Заказ / Партия)"),
        ("H4:H5", "Кол-во\nкомпл."),
        ("I4:K4", "Комплектующие (шт.)"),
        ("L4:L5", "Сдано на\nсклад (шт.)"),
        ("M4:M5", "Причины простоев / Отклонений /\nПримечания"),
    ]

    for cell_range, text in headers:
        if ":" in cell_range:
            ws.merge_cells(cell_range)
            top_left = cell_range.split(":")[0]
            ws[top_left] = text
        else:
            ws[cell_range] = text

    # Подзаголовки комплектующих
    ws["I5"] = "T"
    ws["J5"] = "L"
    ws["K5"] = "Шурупы"

    for r in range(4, 6):
        for c in range(1, 14):
            cell = ws.cell(row=r, column=c)
            cell.font = header_font
            cell.fill = header_fill if r == 4 else sub_header_fill
            cell.alignment = align_center
            cell.border = thin_border

    # Заполнение строк по дням недели
    current_row = 6
    row_num = 1
    total_plan = 0
    total_fact = 0
    total_sets = 0
    total_t = 0
    total_l = 0
    total_screws = 0
    total_warehouse = 0

    curr_d = start_date
    while curr_d <= end_date:
        d_str = curr_d.strftime("%d.%m.%Y")
        day_report = next((rep for rep in reports if rep.report_date == curr_d), None)

        if day_report and day_report.items:
            for item in day_report.items:
                ws.cell(row=current_row, column=1, value=row_num).alignment = align_center
                ws.cell(row=current_row, column=2, value=d_str).alignment = align_center
                ws.cell(row=current_row, column=3, value=day_report.master_name or "Миркасимов И.М.").alignment = align_left
                ws.cell(row=current_row, column=4, value=item.dimension_size or "").alignment = align_left
                ws.cell(row=current_row, column=5, value=item.plan_qty or 0).alignment = align_right
                ws.cell(row=current_row, column=6, value=item.fact_qty or 0).alignment = align_right
                ws.cell(row=current_row, column=7, value=item.order_batch or "").alignment = align_left
                ws.cell(row=current_row, column=8, value=item.sets_count or 0).alignment = align_right
                ws.cell(row=current_row, column=9, value=item.comp_t_count or 0).alignment = align_right
                ws.cell(row=current_row, column=10, value=item.comp_l_count or 0).alignment = align_right
                ws.cell(row=current_row, column=11, value=item.comp_screws_count or 0).alignment = align_right
                
                wh_val = item.warehouse_submitted_qty if item.warehouse_submitted_qty > 0 else (day_report.warehouse_submitted_qty or 0)
                ws.cell(row=current_row, column=12, value=wh_val).alignment = align_right
                
                notes_val = item.notes or day_report.notes or ""
                ws.cell(row=current_row, column=13, value=notes_val).alignment = align_left

                for c in range(1, 14):
                    cell = ws.cell(row=current_row, column=c)
                    cell.font = data_font
                    cell.border = thin_border

                total_plan += (item.plan_qty or 0)
                total_fact += (item.fact_qty or 0)
                total_sets += (item.sets_count or 0)
                total_t += (item.comp_t_count or 0)
                total_l += (item.comp_l_count or 0)
                total_screws += (item.comp_screws_count or 0)
                total_warehouse += wh_val

                current_row += 1
                row_num += 1
        else:
            # Пустая строка для дня
            ws.cell(row=current_row, column=1, value=row_num).alignment = align_center
            ws.cell(row=current_row, column=2, value=d_str).alignment = align_center
            ws.cell(row=current_row, column=3, value=day_report.master_name if day_report else "Миркасимов И.М.").alignment = align_left
            ws.cell(row=current_row, column=4, value="").alignment = align_left
            ws.cell(row=current_row, column=5, value=750).alignment = align_right # Стандартный сменный план
            ws.cell(row=current_row, column=6, value="").alignment = align_right
            ws.cell(row=current_row, column=7, value="").alignment = align_left
            ws.cell(row=current_row, column=8, value="").alignment = align_right
            ws.cell(row=current_row, column=9, value="").alignment = align_right
            ws.cell(row=current_row, column=10, value="").alignment = align_right
            ws.cell(row=current_row, column=11, value="").alignment = align_right
            ws.cell(row=current_row, column=12, value="").alignment = align_right
            ws.cell(row=current_row, column=13, value=day_report.notes if day_report else "").alignment = align_left

            for c in range(1, 14):
                cell = ws.cell(row=current_row, column=c)
                cell.font = data_font
                cell.border = thin_border

            total_plan += 750
            current_row += 1
            row_num += 1

        curr_d += timedelta(days=1)

    # Строка ИТОГО ЗА НЕДЕЛЮ
    ws.merge_cells(f"A{current_row}:D{current_row}")
    ws[f"A{current_row}"] = f"ИТОГО ЗА НЕДЕЛЮ ({start_date.strftime('%d.%m.%Y')} - {end_date.strftime('%d.%m.%Y')}):"
    ws[f"A{current_row}"].font = total_font
    ws[f"A{current_row}"].alignment = align_right

    ws.cell(row=current_row, column=5, value=total_plan).alignment = align_right
    ws.cell(row=current_row, column=6, value=total_fact).alignment = align_right
    ws.cell(row=current_row, column=7, value="")
    ws.cell(row=current_row, column=8, value=total_sets).alignment = align_right
    ws.cell(row=current_row, column=9, value=total_t).alignment = align_right
    ws.cell(row=current_row, column=10, value=total_l).alignment = align_right
    ws.cell(row=current_row, column=11, value=total_screws).alignment = align_right
    ws.cell(row=current_row, column=12, value=total_warehouse).alignment = align_right
    ws.cell(row=current_row, column=13, value="")

    for c in range(1, 14):
        cell = ws.cell(row=current_row, column=c)
        cell.font = total_font
        cell.fill = total_fill
        cell.border = thick_bottom

    # Подписи внизу
    sign_row = current_row + 2
    ws.cell(row=sign_row, column=2, value="Мастер: ____________________ / Миркасимов И.М.").font = bold_font
    ws.cell(row=sign_row, column=5, value="Служба СКК: ____________________ / Мусаилова З.").font = bold_font
    ws.cell(row=sign_row, column=9, value="Начальник цеха / ПТО: ____________________").font = bold_font

    # Ширина колонок
    col_widths = {
        1: 6,   # №
        2: 14,  # Дата
        3: 18,  # Мастер
        4: 26,  # Размер / Типоразмер
        5: 12,  # План
        6: 12,  # Факт
        7: 22,  # Комплект / Заказ
        8: 12,  # Кол-во компл
        9: 10,  # T
        10: 10, # L
        11: 12, # Шурупы
        12: 15, # Сдано на склад
        13: 35  # Примечания
    }
    for col_idx, width in col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Сохраняем в память
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"Рапорт_ЛКЦ_{start_date.strftime('%d.%m')}-{end_date.strftime('%d.%m.%Y')}.xlsx"
    from urllib.parse import quote
    encoded_filename = quote(filename)

    headers = {
        "Content-Disposition": f"attachment; filename*=UTF-8''{encoded_filename}"
    }
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )
